#!/usr/bin/env python3
"""
GDP Production Model generator.

Builds `GDP_Production_Model.xlsx`, a production-approach (output) GDP model:

    Gross Output (commodity)   = Quantity x Price
    GVA (commodity, nominal)   = Gross Output x (1 - IC ratio)
    GVA (commodity, real)      = Quantity x Base-year price x (1 - IC ratio)
    GDP at market prices       = Sum of GVA + taxes less subsidies on products

The user enters production quantities (from commodity boards) and prices in
the blue input cells; everything else is computed with live Excel formulas.

Extension points (edit and re-run this script):
    YEARS            - the year range of the model
    BASE_YEAR        - default constant-price base year (also editable in Settings)
    SECTORS          - commodity-producing sectors (dropdown in Commodities sheet)
    OTHER_SECTORS    - non-commodity industries entered as GVA directly
    SEED_COMMODITIES - example commodity register rows (replace with your own)
    N_COMMODITY_ROWS - number of commodity slots in the register
"""

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------
YEARS = list(range(2019, 2027))          # 2019..2026
BASE_YEAR = 2021
N_COMMODITY_ROWS = 40                    # commodity slots in the register

SECTORS = [
    "Agriculture & Livestock",
    "Forestry",
    "Fishing",
    "Mining & Quarrying",
]

OTHER_SECTORS = [
    "Manufacturing",
    "Electricity & Water",
    "Construction",
    "Wholesale & Retail Trade",
    "Transport & Storage",
    "Accommodation & Food Services",
    "Information & Communication",
    "Financial & Insurance Services",
    "Real Estate",
    "Public Administration",
    "Education & Health",
    "Other Services",
]

# (code, name, sector, unit, IC ratio, source/board) -- EXAMPLE data, replace.
SEED_COMMODITIES = [
    ("AGR01", "Cocoa beans",        SECTORS[0], "tonnes",  0.30, "Cocoa Board"),
    ("AGR02", "Coffee (green)",     SECTORS[0], "tonnes",  0.30, "Coffee Board"),
    ("AGR03", "Rubber (dry)",       SECTORS[0], "tonnes",  0.32, "Rubber Board"),
    ("AGR04", "Palm oil (crude)",   SECTORS[0], "tonnes",  0.35, "Palm Oil Board"),
    ("AGR05", "Rice (paddy)",       SECTORS[0], "tonnes",  0.28, "Grain Board"),
    ("AGR06", "Maize",              SECTORS[0], "tonnes",  0.28, "Grain Board"),
    ("AGR07", "Cotton (seed)",      SECTORS[0], "tonnes",  0.30, "Cotton Board"),
    ("FOR01", "Round logs",         SECTORS[1], "m3",      0.35, "Forestry Authority"),
    ("FOR02", "Sawn timber",        SECTORS[1], "m3",      0.40, "Forestry Authority"),
    ("FSH01", "Fish catch (marine)", SECTORS[2], "tonnes", 0.35, "Fisheries Authority"),
    ("MIN01", "Gold",               SECTORS[3], "kg",      0.40, "Minerals Board"),
    ("MIN02", "Diamonds",           SECTORS[3], "carats",  0.40, "Minerals Board"),
    ("MIN03", "Iron ore",           SECTORS[3], "tonnes",  0.45, "Minerals Board"),
]

# Example quantities / prices so the model demonstrably works out of the box.
# key: code -> {year: (quantity, price)}
def _series(q0, p0, qg, pg):
    """Geometric example series across YEARS."""
    out = {}
    q, p = q0, p0
    for y in YEARS:
        out[y] = (round(q, 1), round(p, 2))
        q *= 1 + qg
        p *= 1 + pg
    return out

SEED_DATA = {
    "AGR01": _series(85_000, 2_100, 0.03, 0.06),
    "AGR02": _series(12_000, 3_400, 0.02, 0.05),
    "AGR03": _series(60_000, 1_500, 0.01, 0.04),
    "AGR04": _series(140_000, 800, 0.04, 0.05),
    "AGR05": _series(260_000, 450, 0.03, 0.04),
    "AGR06": _series(180_000, 320, 0.02, 0.04),
    "AGR07": _series(35_000, 900, 0.01, 0.05),
    "FOR01": _series(400_000, 95, 0.02, 0.03),
    "FOR02": _series(120_000, 260, 0.02, 0.04),
    "FSH01": _series(48_000, 1_900, 0.01, 0.05),
    "MIN01": _series(9_500, 58_000, 0.05, 0.08),
    "MIN02": _series(180_000, 210, 0.02, 0.06),
    "MIN03": _series(2_400_000, 75, 0.03, 0.05),
}

# ----------------------------------------------------------------------------
# Layout constants
# ----------------------------------------------------------------------------
NY = len(YEARS)
FIRST_YEAR_COL = 3                       # column C on every data sheet
LAST_YEAR_COL = FIRST_YEAR_COL + NY - 1
YEAR_COLS = [get_column_letter(c) for c in range(FIRST_YEAR_COL, LAST_YEAR_COL + 1)]
DATA_FIRST_ROW = 2                       # commodity sheets: data starts on row 2
DATA_LAST_ROW = DATA_FIRST_ROW + N_COMMODITY_ROWS - 1

# ----------------------------------------------------------------------------
# Styles
# ----------------------------------------------------------------------------
F_TITLE = Font(bold=True, size=14, color="FF1F3864")
F_HEADER = Font(bold=True, color="FFFFFFFF")
F_BOLD = Font(bold=True)
F_NOTE = Font(italic=True, size=9, color="FF7F7F7F")
FILL_HEADER = PatternFill("solid", fgColor="FF1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFDDEBF7")     # light blue = INPUT
FILL_CALC = PatternFill("solid", fgColor="FFF2F2F2")      # light grey = CALC
FILL_TOTAL = PatternFill("solid", fgColor="FFFFF2CC")     # amber = key result
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NUM_QTY = "#,##0"
NUM_PRICE = "#,##0.00"
NUM_VAL = "#,##0"
NUM_PCT = "0.0%"
NUM_IDX = "0.0"


def style_header_row(ws, row, last_col):
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = F_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def box(ws, row, col, value=None, formula=None, fill=None, numfmt=None,
        bold=False, fmt_font=None):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    elif value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if numfmt is not None:
        cell.number_format = numfmt
    if bold:
        cell.font = F_BOLD
    if fmt_font is not None:
        cell.font = fmt_font
    cell.border = BORDER
    return cell


# ----------------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------------
wb = Workbook()

# ---------------------------------------------------------------- READ_ME ---
ws = wb.active
ws.title = "READ_ME"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110
rows = [
    ("GDP PRODUCTION MODEL", F_TITLE),
    ("Production (output) approach:  GDP = Gross Value Added + Taxes less subsidies on products", F_BOLD),
    ("", None),
    ("HOW TO USE", F_BOLD),
    ("1. Settings — set your country, currency unit, first year and constant-price base year.", None),
    ("2. Commodities — register each commodity: name, sector, unit, IC ratio and the commodity board it comes from.", None),
    ("   The IC ratio is the share of gross output used up as intermediate consumption (inputs). GVA = Output x (1 - IC ratio).", None),
    ("3. Production — enter yearly production quantities from your commodity boards (blue cells).", None),
    ("4. Prices — enter yearly average prices per unit, in your currency, same unit as the quantity (blue cells).", None),
    ("5. Other_Sectors — enter GVA directly for industries without commodity detail (manufacturing, services, government...).", None),
    ("6. GDP_Summary — GDP appears automatically: nominal, real (constant prices), growth, deflator and per capita.", None),
    ("   Enter 'Taxes less subsidies on products' and 'Population' on that sheet to complete the picture.", None),
    ("", None),
    ("COLOUR CODE", F_BOLD),
    ("   BLUE cells   = your inputs", None),
    ("   GREY cells   = calculated, do not type over them", None),
    ("   AMBER cells  = key results", None),
    ("", None),
    ("ADDING A COMMODITY", F_BOLD),
    ("   Add it on any empty row of the Commodities sheet — its name flows automatically to Production, Prices and the", None),
    ("   GVA sheets. Then type its quantities and prices. Totals and GDP update by themselves (40 slots are pre-wired).", None),
    ("", None),
    ("REAL (CONSTANT-PRICE) GDP", F_BOLD),
    ("   Real GVA values every year's quantity at the base-year price set in Settings, isolating volume growth from", None),
    ("   price changes. Make sure prices are filled in for the base year.", None),
    ("", None),
    ("NOTE", F_BOLD),
    ("   All values are in the currency unit you choose (e.g. millions). Keep quantity units and price units consistent:", None),
    ("   if quantity is in tonnes, the price must be per tonne. The data shipped in this file is EXAMPLE data — replace it.", None),
]
for i, (text, font) in enumerate(rows, start=2):
    c = ws.cell(row=i, column=2, value=text)
    if font:
        c.font = font

# ---------------------------------------------------------------- Settings --
ws = wb.create_sheet("Settings")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 22
ws.column_dimensions["D"].width = 80
ws["A1"] = "SETTINGS"
ws["A1"].font = F_TITLE
settings = [
    ("Country / economy", "Example Country", None),
    ("Base year (constant prices)", BASE_YEAR, "Must be a year with prices filled in on the Prices sheet."),
    ("First year of the model", YEARS[0], f"The model covers {NY} consecutive years from this year."),
    ("Currency unit", "LCU", "e.g. 'GHS million' — label only, used in headings."),
]
for i, (label, value, note) in enumerate(settings, start=3):
    ws.cell(row=i, column=1, value=label).font = F_BOLD
    box(ws, i, 2, value=value, fill=FILL_INPUT)
    if note:
        ws.cell(row=i, column=4, value=note).font = F_NOTE
# Settings!B4 = base year, B5 = first year, B6 = currency (rows 3..6)

# ------------------------------------------------------------- Commodities --
ws = wb.create_sheet("Commodities")
headers = ["Code", "Commodity", "Sector", "Unit", "IC ratio", "Source (board)", "Notes"]
widths = [10, 26, 24, 12, 10, 22, 30]
for c, (h, w) in enumerate(zip(headers, widths), start=1):
    ws.cell(row=1, column=c, value=h)
    ws.column_dimensions[get_column_letter(c)].width = w
style_header_row(ws, 1, len(headers))
ws.freeze_panes = "A2"

for i in range(N_COMMODITY_ROWS):
    r = DATA_FIRST_ROW + i
    seed = SEED_COMMODITIES[i] if i < len(SEED_COMMODITIES) else None
    for c in range(1, len(headers) + 1):
        val = seed[c - 1] if seed and c <= 6 else None
        cell = box(ws, r, c, value=val, fill=FILL_INPUT)
        if c == 5:
            cell.number_format = "0%"

dv_sector = DataValidation(
    type="list",
    formula1='"' + ",".join(SECTORS) + '"',
    allow_blank=True,
    showErrorMessage=True,
    errorTitle="Unknown sector",
    error="Pick one of the commodity sectors (or add it to SECTORS in build_model.py and re-run).",
)
ws.add_data_validation(dv_sector)
dv_sector.add(f"C{DATA_FIRST_ROW}:C{DATA_LAST_ROW}")

dv_ic = DataValidation(
    type="decimal", operator="between", formula1="0", formula2="0.95",
    allow_blank=True, showErrorMessage=True,
    errorTitle="IC ratio", error="Enter a share between 0 and 0.95 (e.g. 0.30 = 30% of output used as inputs).",
)
ws.add_data_validation(dv_ic)
dv_ic.add(f"E{DATA_FIRST_ROW}:E{DATA_LAST_ROW}")


def year_header(ws, row=1, source=None):
    """Write linked year headers into C..J of `row`."""
    for j, col in enumerate(YEAR_COLS):
        if source:
            f = f"={source}!{col}${1}"
        elif j == 0:
            f = "=Settings!$B$5"
        else:
            f = f"={YEAR_COLS[j-1]}{row}+1"
        ws[f"{col}{row}"] = f


def commodity_data_sheet(name, input_sheet, numfmt):
    """Production / Prices: A=Commodity, B=Unit (linked), C..=years (input)."""
    ws = wb.create_sheet(name)
    ws.cell(row=1, column=1, value="Commodity")
    ws.cell(row=1, column=2, value="Unit")
    year_header(ws, 1, source=None if input_sheet else "Production")
    if name != "Production":
        for col in YEAR_COLS:
            ws[f"{col}1"] = f"=Production!{col}$1"
    style_header_row(ws, 1, LAST_YEAR_COL)
    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    for col in YEAR_COLS:
        ws.column_dimensions[col].width = 13
    for i in range(N_COMMODITY_ROWS):
        r = DATA_FIRST_ROW + i
        box(ws, r, 1, formula=f'=IF(Commodities!B{r}="","",Commodities!B{r})', fill=FILL_CALC)
        box(ws, r, 2, formula=f'=IF(Commodities!D{r}="","",Commodities!D{r})', fill=FILL_CALC)
        for col in YEAR_COLS:
            box(ws, r, ws[f"{col}1"].column, fill=FILL_INPUT, numfmt=numfmt)
    return ws


ws_prod = commodity_data_sheet("Production", True, NUM_QTY)
ws_price = commodity_data_sheet("Prices", True, NUM_PRICE)
ws_prod["A1"].comment = None

# seed example data
code_row = {SEED_COMMODITIES[i][0]: DATA_FIRST_ROW + i for i in range(len(SEED_COMMODITIES))}
for code, series in SEED_DATA.items():
    r = code_row[code]
    for j, y in enumerate(YEARS):
        q, p = series[y]
        ws_prod[f"{YEAR_COLS[j]}{r}"] = q
        ws_price[f"{YEAR_COLS[j]}{r}"] = p

# ------------------------------------------------------------ GVA sheets ----
def gva_sheet(name, real):
    ws = wb.create_sheet(name)
    ws.cell(row=1, column=1, value="Commodity")
    ws.cell(row=1, column=2, value="Sector")
    for col in YEAR_COLS:
        ws[f"{col}1"] = f"=Production!{col}$1"
    style_header_row(ws, 1, LAST_YEAR_COL)
    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 24
    for col in YEAR_COLS:
        ws.column_dimensions[col].width = 14
    base_col = None
    if real:
        base_col = get_column_letter(LAST_YEAR_COL + 2)
        ws[f"{base_col}1"] = "Base-year price"
        ws.column_dimensions[base_col].width = 16
        style_header_row(ws, 1, LAST_YEAR_COL)
        h = ws[f"{base_col}1"]
        h.font = F_HEADER; h.fill = FILL_HEADER; h.border = BORDER
        h.alignment = Alignment(horizontal="center")

    for i in range(N_COMMODITY_ROWS):
        r = DATA_FIRST_ROW + i
        box(ws, r, 1, formula=f'=IF(Commodities!B{r}="","",Commodities!B{r})', fill=FILL_CALC)
        box(ws, r, 2, formula=f'=IF(Commodities!C{r}="","",Commodities!C{r})', fill=FILL_CALC)
        if real:
            f_base = (
                f'=IF($A{r}="","",IFERROR(INDEX(Prices!${YEAR_COLS[0]}{r}:${YEAR_COLS[-1]}{r},'
                f'MATCH(Settings!$B$4,Prices!${YEAR_COLS[0]}$1:${YEAR_COLS[-1]}$1,0)),""))'
            )
            box(ws, r, LAST_YEAR_COL + 2, formula=f_base, fill=FILL_CALC, numfmt=NUM_PRICE)
        for col in YEAR_COLS:
            if real:
                f = (
                    f'=IF(OR($A{r}="",Production!{col}{r}="",${base_col}{r}=""),"",'
                    f"Production!{col}{r}*${base_col}{r}*(1-Commodities!$E{r}))"
                )
            else:
                f = (
                    f'=IF(OR($A{r}="",Production!{col}{r}="",Prices!{col}{r}=""),"",'
                    f"Production!{col}{r}*Prices!{col}{r}*(1-Commodities!$E{r}))"
                )
            box(ws, r, ws[f"{col}1"].column, formula=f, fill=FILL_CALC, numfmt=NUM_VAL)

    tr = DATA_LAST_ROW + 2
    box(ws, tr, 1, value="TOTAL commodity GVA", bold=True, fill=FILL_TOTAL)
    box(ws, tr, 2, fill=FILL_TOTAL)
    for col in YEAR_COLS:
        box(ws, tr, ws[f"{col}1"].column,
            formula=f"=SUM({col}{DATA_FIRST_ROW}:{col}{DATA_LAST_ROW})",
            bold=True, fill=FILL_TOTAL, numfmt=NUM_VAL)
    return ws


ws_gva_n = gva_sheet("GVA_Nominal", real=False)
ws_gva_r = gva_sheet("GVA_Real", real=True)

# ---------------------------------------------------------- Other sectors ---
ws = wb.create_sheet("Other_Sectors")
ws.column_dimensions["A"].width = 32
for col in YEAR_COLS:
    ws.column_dimensions[col].width = 14
ws["A1"] = "OTHER INDUSTRIES — GVA entered directly (no commodity detail)"
ws["A1"].font = F_TITLE

OTHER_N_HDR = 3                       # nominal block header row
OTHER_N_FIRST = OTHER_N_HDR + 1       # rows 4..15
OTHER_N_LAST = OTHER_N_FIRST + len(OTHER_SECTORS) - 1
OTHER_R_HDR = OTHER_N_LAST + 3        # real block header row
OTHER_R_FIRST = OTHER_R_HDR + 1
OTHER_R_LAST = OTHER_R_FIRST + len(OTHER_SECTORS) - 1

ws.cell(row=OTHER_N_HDR, column=1, value="Nominal GVA (current prices)")
ws.cell(row=OTHER_R_HDR, column=1, value="Real GVA (constant base-year prices)")
for hdr in (OTHER_N_HDR, OTHER_R_HDR):
    for col in YEAR_COLS:
        ws[f"{col}{hdr}"] = f"=Production!{col}$1"
    style_header_row(ws, hdr, LAST_YEAR_COL)

for blk_first, linked in ((OTHER_N_FIRST, False), (OTHER_R_FIRST, True)):
    for i, sector in enumerate(OTHER_SECTORS):
        r = blk_first + i
        if linked:
            box(ws, r, 1, formula=f"=A{OTHER_N_FIRST + i}", fill=FILL_CALC, bold=False)
        else:
            box(ws, r, 1, value=sector, fill=FILL_INPUT)
        for col in YEAR_COLS:
            box(ws, r, ws[f"{col}{OTHER_N_HDR}"].column, fill=FILL_INPUT, numfmt=NUM_VAL)
ws.freeze_panes = "B2"

# ------------------------------------------------------------ GDP summary ---
ws = wb.create_sheet("GDP_Summary")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 40
for col in YEAR_COLS:
    ws.column_dimensions[col].width = 14

ws["A1"] = "GDP — PRODUCTION APPROACH"
ws["A1"].font = F_TITLE
ws["A2"] = '=CONCATENATE("Values in ",Settings!B6,".  Constant prices of base year ",Settings!B4,".")'
ws["A2"].font = F_NOTE

YEAR_HDR = 3
for col in YEAR_COLS:
    ws[f"{col}{YEAR_HDR}"] = f"=Production!{col}$1"
ws.cell(row=YEAR_HDR, column=1, value="Industry / aggregate")
style_header_row(ws, YEAR_HDR, LAST_YEAR_COL)
ws.freeze_panes = f"{YEAR_COLS[0]}{YEAR_HDR + 1}"

r = YEAR_HDR + 1


def section(label):
    global r
    ws.cell(row=r, column=1, value=label).font = F_BOLD
    ws.cell(row=r, column=1).fill = FILL_TOTAL
    for col in YEAR_COLS:
        ws[f"{col}{r}"].fill = FILL_TOTAL
    r += 1


def blank():
    global r
    r += 1


def row_of(label, per_col_formula=None, fill=FILL_CALC, numfmt=NUM_VAL,
           bold=False, input_row=False, indent=False):
    """Add one line; per_col_formula(col_letter) -> formula string.

    Indentation must be via cell formatting, never leading spaces: the
    label cell doubles as the SUMIFS criteria against sector names.
    """
    global r
    cell = box(ws, r, 1, value=label, bold=bold,
               fill=FILL_TOTAL if bold else None)
    if indent:
        cell.alignment = Alignment(indent=1)
    for col in YEAR_COLS:
        f = per_col_formula(col) if per_col_formula else None
        box(ws, r, ws[f"{col}{YEAR_HDR}"].column, formula=f,
            fill=FILL_INPUT if input_row else (FILL_TOTAL if bold else fill),
            numfmt=numfmt, bold=bold)
    this = r
    r += 1
    return this


gva_rows = {}

section("NOMINAL — current prices")
nom_first = r
for s in SECTORS:
    gva_rows[("N", s)] = row_of(
        s,
        lambda col, s=s: (
            f'=SUMIFS(GVA_Nominal!{col}${DATA_FIRST_ROW}:{col}${DATA_LAST_ROW},'
            f'GVA_Nominal!$B${DATA_FIRST_ROW}:$B${DATA_LAST_ROW},$A{r})'
        ),
    )
for i, s in enumerate(OTHER_SECTORS):
    gva_rows[("N", s)] = row_of(
        s,
        lambda col, i=i: f"=IF(Other_Sectors!{col}{OTHER_N_FIRST + i}=\"\",0,Other_Sectors!{col}{OTHER_N_FIRST + i})",
    )
nom_last = r - 1
r_gva_n = row_of(
    "Gross Value Added at basic prices",
    lambda col: f"=SUM({col}{nom_first}:{col}{nom_last})",
    bold=True,
)
r_tax_n = row_of("Taxes less subsidies on products  [input]", input_row=True)
r_gdp_n = row_of(
    "GDP at market prices — NOMINAL",
    lambda col: f'=IF({col}{r_gva_n}=0,"",{col}{r_gva_n}+IF({col}{r_tax_n}="",0,{col}{r_tax_n}))',
    bold=True,
)
blank()

section("REAL — constant base-year prices")
real_first = r
for s in SECTORS:
    gva_rows[("R", s)] = row_of(
        s,
        lambda col, s=s: (
            f'=SUMIFS(GVA_Real!{col}${DATA_FIRST_ROW}:{col}${DATA_LAST_ROW},'
            f'GVA_Real!$B${DATA_FIRST_ROW}:$B${DATA_LAST_ROW},$A{r})'
        ),
    )
for i, s in enumerate(OTHER_SECTORS):
    gva_rows[("R", s)] = row_of(
        s,
        lambda col, i=i: f"=IF(Other_Sectors!{col}{OTHER_R_FIRST + i}=\"\",0,Other_Sectors!{col}{OTHER_R_FIRST + i})",
    )
real_last = r - 1
r_gva_r = row_of(
    "Gross Value Added at basic prices (real)",
    lambda col: f"=SUM({col}{real_first}:{col}{real_last})",
    bold=True,
)
r_tax_r = row_of("Taxes less subsidies, constant prices  [input]", input_row=True)
r_gdp_r = row_of(
    "GDP at constant prices — REAL",
    lambda col: f'=IF({col}{r_gva_r}=0,"",{col}{r_gva_r}+IF({col}{r_tax_r}="",0,{col}{r_tax_r}))',
    bold=True,
)
blank()

section("INDICATORS")


def growth_formula(target_row):
    def f(col):
        idx = YEAR_COLS.index(col)
        if idx == 0:
            return None
        prev = YEAR_COLS[idx - 1]
        return (
            f'=IF(OR({prev}{target_row}="",{col}{target_row}="",{prev}{target_row}=0),"",'
            f"{col}{target_row}/{prev}{target_row}-1)"
        )
    return f


r_growth_r = row_of("Real GDP growth (%)", growth_formula(r_gdp_r), numfmt=NUM_PCT)
r_growth_n = row_of("Nominal GDP growth (%)", growth_formula(r_gdp_n), numfmt=NUM_PCT)
r_defl = row_of(
    "GDP deflator (base year = 100)",
    lambda col: f'=IF(OR({col}{r_gdp_r}="",{col}{r_gdp_r}=0,{col}{r_gdp_n}=""),"",{col}{r_gdp_n}/{col}{r_gdp_r}*100)',
    numfmt=NUM_IDX,
)
r_pop = row_of("Population  [input]", input_row=True, numfmt=NUM_QTY)
r_pc = row_of(
    "Nominal GDP per capita",
    lambda col: f'=IF(OR({col}{r_pop}="",{col}{r_pop}=0,{col}{r_gdp_n}=""),"",{col}{r_gdp_n}/{col}{r_pop})',
    numfmt=NUM_PRICE,
)
blank()

section("STRUCTURE — share of nominal GVA")
share_first = r
for s in SECTORS + OTHER_SECTORS:
    src = gva_rows[("N", s)]
    row_of(
        s,
        lambda col, src=src: f'=IF({col}{r_gva_n}=0,"",{col}{src}/{col}{r_gva_n})',
        numfmt=NUM_PCT,
    )
share_last = r - 1

# ------------------------------------------------------------- Dashboard ----
ws_dash = wb.create_sheet("Dashboard")
ws_dash.sheet_view.showGridLines = False
ws_dash["A1"] = "DASHBOARD"
ws_dash["A1"].font = F_TITLE

sumry = wb["GDP_Summary"]
cats = Reference(sumry, min_col=FIRST_YEAR_COL, max_col=LAST_YEAR_COL, min_row=YEAR_HDR)

line = LineChart()
line.title = "GDP: nominal vs real"
line.height, line.width = 9, 18
for row_i, name in ((r_gdp_n, "Nominal GDP"), (r_gdp_r, "Real GDP")):
    ref = Reference(sumry, min_col=FIRST_YEAR_COL, max_col=LAST_YEAR_COL, min_row=row_i)
    s = Series(ref, title=name)
    line.series.append(s)
line.set_categories(cats)
line.y_axis.title = "Value"
ws_dash.add_chart(line, "A3")

bar = BarChart()
bar.title = "Real GDP growth"
bar.height, bar.width = 9, 18
ref = Reference(sumry, min_col=FIRST_YEAR_COL, max_col=LAST_YEAR_COL, min_row=r_growth_r)
bar.series.append(Series(ref, title="Real growth"))
bar.set_categories(cats)
ws_dash.add_chart(bar, "A22")

pie = PieChart()
pie.title = f"GVA composition, {YEARS[-1]} (nominal shares)"
pie.height, pie.width = 11, 18
data = Reference(sumry, min_col=LAST_YEAR_COL, max_col=LAST_YEAR_COL,
                 min_row=share_first, max_row=share_last)
labels = Reference(sumry, min_col=1, max_col=1, min_row=share_first, max_row=share_last)
pie.add_data(data, titles_from_data=False)
pie.set_categories(labels)
ws_dash.add_chart(pie, "L3")

# ---------------------------------------------------------------- finish ----
order = ["READ_ME", "Settings", "Commodities", "Production", "Prices",
         "Other_Sectors", "GVA_Nominal", "GVA_Real", "GDP_Summary", "Dashboard"]
wb._sheets = [wb[name] for name in order]

OUT = "GDP_Production_Model.xlsx"
wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Summary rows: nominal GDP={r_gdp_n}, real GDP={r_gdp_r}, growth={r_growth_r}")
